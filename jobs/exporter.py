import csv
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from jobs.models import Job, Company


HEADERS = [
    "ID",
    "Job ID",
    "Title",
    "Company",
    "URL",
    "Location",
    "Description",
    "Employment type",
    "Seniority",
    "Details status",
    "Attempts",
    "Last error",
    "Published at",
    "Created at",
    "Updated at",
]

COMPANY_HEADERS = [
    "ID",
    "Name",
    "LinkedIn URL",
    "Description",
    "Industry",
    "Website",
    "Employee count",
    "Created at",
    "Updated at",
]


def excel_value(value):
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return value


def get_jobs_data():
    jobs = Job.objects.select_related("company").order_by("id")

    for job in jobs.iterator(chunk_size=1000):
        company = job.company

        yield [
            job.id,
            job.job_id,
            job.title,
            company.name if company else "",
            job.job_url,
            job.location,
            job.description,
            job.employment_type,
            job.seniority,
            job.details_status,
            job.attempts,
            job.last_error,
            job.published_at,
            job.created_at,
            job.updated_at,
        ]


def get_companies_data():
    companies = Company.objects.order_by("id")

    for company in companies.iterator(chunk_size=1000):
        yield [
            company.id,
            company.name,
            company.linkedin_url,
            company.description,
            company.industry,
            company.website,
            company.employee_count,
            company.created_at,
            company.updated_at,
        ]


def export_jobs(file_path):
    count = 0

    with open(file_path, "w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file)
        writer.writerow(HEADERS)

        for row in get_jobs_data():
            writer.writerow(row)
            count += 1

    return count


def export_companies(file_path):
    count = 0

    with open(file_path, "w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file)
        writer.writerow(COMPANY_HEADERS)

        for row in get_companies_data():
            writer.writerow(row)
            count += 1

    return count


def style_header(sheet, color="1F4E78"):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=color,
        )

    sheet.freeze_panes = "A2"


def auto_adjust_columns(sheet):
    for column_index, column_cells in enumerate(
        sheet.iter_cols(), start=1
    ):
        max_length = max(
            (
                len(str(cell.value or ""))
                for cell in list(column_cells)[:200]
            ),
            default=12,
        )

        sheet.column_dimensions[
            get_column_letter(column_index)
        ].width = min(max(max_length + 2, 12), 45)


def export_jobs_excel(file_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Jobs"

    sheet.append(HEADERS)
    style_header(sheet)

    count = 0

    for row in get_jobs_data():
        sheet.append([excel_value(value) for value in row])
        count += 1

    sheet.auto_filter.ref = sheet.dimensions
    auto_adjust_columns(sheet)

    workbook.save(file_path)
    return count


def export_all_excel(file_path):
    workbook = Workbook()

    # Аркуш вакансій із даними компаній
    jobs_sheet = workbook.active
    jobs_sheet.title = "Jobs"

    jobs_headers = HEADERS + [
        "Company LinkedIn URL",
        "Company Description",
        "Company Industry",
        "Company Website",
        "Company Employee Count",
    ]

    jobs_sheet.append(jobs_headers)
    style_header(jobs_sheet)

    jobs = Job.objects.select_related("company").order_by("id")
    jobs_count = 0

    for job in jobs.iterator(chunk_size=1000):
        company = job.company

        row = [
            job.id,
            job.job_id,
            job.title,
            company.name if company else "",
            job.job_url,
            job.location,
            job.description,
            job.employment_type,
            job.seniority,
            job.details_status,
            job.attempts,
            job.last_error,
            job.published_at,
            job.created_at,
            job.updated_at,
            company.linkedin_url if company else "",
            company.description if company else "",
            company.industry if company else "",
            company.website if company else "",
            company.employee_count if company else "",
        ]

        jobs_sheet.append([excel_value(value) for value in row])
        jobs_count += 1

    jobs_sheet.auto_filter.ref = jobs_sheet.dimensions
    auto_adjust_columns(jobs_sheet)

    # Окремий аркуш компаній
    companies_sheet = workbook.create_sheet("Companies")
    companies_sheet.append(COMPANY_HEADERS)
    style_header(companies_sheet, color="38761D")

    companies_count = 0

    for row in get_companies_data():
        companies_sheet.append(
            [excel_value(value) for value in row]
        )
        companies_count += 1

    companies_sheet.auto_filter.ref = companies_sheet.dimensions
    auto_adjust_columns(companies_sheet)

    workbook.save(file_path)

    return jobs_count, companies_count